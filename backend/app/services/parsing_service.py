from dataclasses import dataclass
from html.parser import HTMLParser
from io import BytesIO, StringIO
from pathlib import Path
import tempfile
import zipfile

from pypdf import PdfReader

from app.ai.ocr_provider import get_ocr_provider
from app.ai.speech_provider import get_speech_to_text_provider
from app.core.config import settings
from app.models.document import Document, DocumentSourceType
from app.services.job_error_service import PublicJobError
from app.storage.storage_service import (
    StorageBackendChangedError,
    materialize_for_backend,
)


class DocumentParsingError(PublicJobError):
    pass


# 单个 PDF 页面抽取出的可见文本低于该阈值时视为扫描版页面。
SCANNED_PDF_PAGE_MIN_CHARS = 16


@dataclass(frozen=True)
class ParsedDocument:
    raw_text: str
    metadata: dict


def parse_document(document: Document) -> ParsedDocument:
    if document.source_type == DocumentSourceType.NOTE.value:
        return ParsedDocument(raw_text=document.raw_text or "", metadata=document.metadata_ or {})

    if not document.file_path:
        raise DocumentParsingError("资料缺少文件路径")
    if not document.storage_backend or not document.storage_scope:
        raise DocumentParsingError("资料缺少存储位置信息")

    try:
        with materialize_for_backend(
            document.file_path,
            document.storage_backend,
            document.storage_scope,
        ) as path:
            if document.source_type == DocumentSourceType.TXT.value:
                return _parse_text_file(path)
            if document.source_type == DocumentSourceType.MARKDOWN.value:
                return _parse_markdown_file(path)
            if document.source_type == DocumentSourceType.PDF.value:
                return _parse_pdf_file(path)
            if document.source_type == DocumentSourceType.IMAGE.value:
                return _parse_image_file(path, document.mime_type)
            if document.source_type == DocumentSourceType.AUDIO.value:
                return _parse_audio_file(path, document.mime_type)
            if document.source_type == DocumentSourceType.DOCX.value:
                return _parse_docx_file(path)
            if document.source_type == DocumentSourceType.HTML.value:
                return _parse_html_file(path)
            if document.source_type == DocumentSourceType.EXCEL.value:
                return _parse_excel_file(path)
            if document.source_type == DocumentSourceType.EPUB.value:
                return _parse_epub_file(path)
            raise DocumentParsingError("暂不支持的资料类型")
    except (DocumentParsingError, StorageBackendChangedError):
        raise
    except Exception as exc:
        raise DocumentParsingError(
            _parsing_failure_message(document.source_type)
        ) from exc


def _parse_text_file(path: Path) -> ParsedDocument:
    text = _read_text(path)
    return ParsedDocument(raw_text=text, metadata={"parser": "text"})


def _parse_markdown_file(path: Path) -> ParsedDocument:
    text = _read_text(path)
    return ParsedDocument(raw_text=text, metadata={"parser": "markdown"})


def _parse_pdf_file(path: Path) -> ParsedDocument:
    reader = PdfReader(str(path))
    page_count = len(reader.pages)
    page_texts = [(page.extract_text() or "").strip() for page in reader.pages]
    low_text_indexes = [
        index
        for index, text in enumerate(page_texts)
        if len(text) < SCANNED_PDF_PAGE_MIN_CHARS
    ]

    ocr_used = False
    if low_text_indexes:
        if settings.ocr_provider == "disabled":
            if not any(page_texts):
                raise DocumentParsingError(
                    "这是一个扫描版 PDF，没有可提取的文字。请配置 OCR Provider 后重试。"
                )
            # 混合型 PDF：低文字页（如封面图片）保持原样，其余页照常入索引。
        else:
            if len(low_text_indexes) > SCANNED_PDF_OCR_PAGE_LIMIT:
                raise DocumentParsingError(
                    f"PDF 中有 {len(low_text_indexes)} 页没有足够的可提取文字，"
                    f"超过单次 OCR 上限（{SCANNED_PDF_OCR_PAGE_LIMIT} 页），请拆分后重新上传"
                )
            ocr_texts = _ocr_pdf_pages(path, low_text_indexes)
            for index, text in zip(low_text_indexes, ocr_texts):
                page_texts[index] = text
            ocr_used = True

    output = StringIO()
    wrote_page = False
    for index, text in enumerate(page_texts, start=1):
        if text:
            if wrote_page:
                output.write("\n\n\n")
            output.write(f"[第 {index} 页]\n{text}")
            wrote_page = True

    return ParsedDocument(
        raw_text=output.getvalue(),
        metadata={
            "parser": "pypdf+ocr" if ocr_used else "pypdf",
            "page_count": page_count,
        },
    )


# 单个 PDF 允许回退 OCR 的最大页数，防止超大扫描件产生失控的 OCR 开销。
SCANNED_PDF_OCR_PAGE_LIMIT = 100


def _ocr_pdf_pages(path: Path, page_indexes: list[int]) -> list[str]:
    """对指定页（0 基）渲染图片并 OCR，返回与 page_indexes 一一对应的文本。"""
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover - 依赖在 requirements 中声明
        raise DocumentParsingError(
            "识别扫描版 PDF 需要 OCR Provider，且后端需安装 PyMuPDF。"
        ) from exc

    ocr = get_ocr_provider()
    results: list[str] = []
    with fitz.open(path) as doc:
        for index in page_indexes:
            pix = doc[index].get_pixmap(dpi=150)
            png = pix.tobytes("png")
            with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as handle:
                handle.write(png)
                temp_path = Path(handle.name)
            try:
                results.append(ocr.extract_text(temp_path, mime_type="image/png").strip())
            finally:
                temp_path.unlink(missing_ok=True)
    if page_indexes and not any(results):
        raise DocumentParsingError("扫描页识别失败，未得到任何文字，请检查 OCR Provider 配置")
    return results


def _parse_image_file(path: Path, mime_type: str | None) -> ParsedDocument:
    text = get_ocr_provider().extract_text(path, mime_type=mime_type)
    return ParsedDocument(raw_text=text, metadata={"parser": "ocr"})


def _parse_audio_file(path: Path, mime_type: str | None) -> ParsedDocument:
    text = get_speech_to_text_provider().transcribe(path, mime_type=mime_type)
    return ParsedDocument(raw_text=text, metadata={"parser": "speech_to_text"})


def _parse_docx_file(path: Path) -> ParsedDocument:
    from docx import Document as DocxDocument

    doc = DocxDocument(str(path))
    parts: list[str] = []
    for paragraph in doc.paragraphs:
        text = paragraph.text.strip()
        if text:
            parts.append(text)
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = "\t".join(cell for cell in cells if cell)
            if line:
                parts.append(line)
    return ParsedDocument(raw_text="\n\n".join(parts), metadata={"parser": "docx"})


def _parse_html_file(path: Path) -> ParsedDocument:
    text = _html_to_text(_read_text(path))
    return ParsedDocument(raw_text=text, metadata={"parser": "html"})


class _HTMLTextExtractor(HTMLParser):
    _SKIP_TAGS = {"script", "style", "noscript", "template", "head"}
    _BLOCK_TAGS = {
        "p", "div", "section", "article", "li", "tr", "h1", "h2", "h3",
        "h4", "h5", "h6", "br", "hr", "blockquote", "pre", "table",
        "ul", "ol", "header", "footer", "nav",
    }

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip_depth = 0
        self._chunks: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in self._SKIP_TAGS:
            self._skip_depth += 1
        elif tag in self._BLOCK_TAGS:
            self._chunks.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in self._SKIP_TAGS and self._skip_depth > 0:
            self._skip_depth -= 1
        elif tag in self._BLOCK_TAGS:
            self._chunks.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip_depth == 0 and data.strip():
            self._chunks.append(data.strip())

    def get_text(self) -> str:
        lines: list[str] = []
        for chunk in "".join(self._chunks).split("\n"):
            collapsed = " ".join(chunk.split())
            if collapsed:
                lines.append(collapsed)
        return "\n\n".join(lines)


def _html_to_text(html: str) -> str:
    extractor = _HTMLTextExtractor()
    extractor.feed(html)
    return extractor.get_text()


def _parse_excel_file(path: Path) -> ParsedDocument:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    output = StringIO()
    try:
        for worksheet in workbook.worksheets:
            output.write(f"## {worksheet.title}\n")
            for row in worksheet.iter_rows(values_only=True):
                cells = ["" if value is None else str(value).strip() for value in row]
                if any(cells):
                    output.write("\t".join(cells) + "\n")
            output.write("\n")
    finally:
        workbook.close()
    return ParsedDocument(raw_text=output.getvalue(), metadata={"parser": "openpyxl"})


def _parse_epub_file(path: Path) -> ParsedDocument:
    from xml.etree import ElementTree

    container_ns = {"cn": "urn:oasis:names:tc:opendocument:xmlns:container"}
    opf_ns = {"opf": "http://www.idpf.org/2007/opf"}
    with zipfile.ZipFile(path) as archive:
        container = archive.read("META-INF/container.xml")
        root_file = ElementTree.fromstring(container).find(
            "cn:rootfiles/cn:rootfile", container_ns
        )
        if root_file is None:
            raise ValueError("EPUB 缺少 OPF 描述文件")
        opf_path = root_file.attrib["full-path"]
        opf_dir = opf_path.rsplit("/", 1)[0] if "/" in opf_path else ""
        opf = ElementTree.fromstring(archive.read(opf_path))

        manifest: dict[str, str] = {}
        for item in opf.findall("opf:manifest/opf:item", opf_ns):
            manifest[item.attrib["id"]] = item.attrib["href"]
        spine_ids = [
            itemref.attrib["idref"]
            for itemref in opf.findall("opf:spine/opf:itemref", opf_ns)
        ]

        parts: list[str] = []
        for spine_id in spine_ids:
            href = manifest.get(spine_id)
            if not href:
                continue
            member_path = f"{opf_dir}/{href}" if opf_dir else href
            try:
                content = archive.read(member_path).decode("utf-8", errors="ignore")
            except KeyError:
                continue
            text = _html_to_text(content)
            if text.strip():
                parts.append(text)
    return ParsedDocument(raw_text="\n\n".join(parts), metadata={"parser": "epub"})


def _read_text(path: Path) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return path.read_text(encoding=encoding)
        except UnicodeDecodeError:
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def _parsing_failure_message(source_type: str) -> str:
    if source_type in {
        DocumentSourceType.TXT.value,
        DocumentSourceType.MARKDOWN.value,
    }:
        return "文本资料读取失败，请确认文件仍然存在且内容可读"
    if source_type == DocumentSourceType.PDF.value:
        return "PDF 解析失败，请确认文件完整且格式有效"
    if source_type == DocumentSourceType.IMAGE.value:
        return "图片文字识别失败，请检查 OCR Provider 配置后重试"
    if source_type == DocumentSourceType.AUDIO.value:
        return "音频转写失败，请检查语音转文字 Provider 配置后重试"
    if source_type == DocumentSourceType.DOCX.value:
        return "Word 文档解析失败，请确认文件完整且为 .docx 格式"
    if source_type == DocumentSourceType.HTML.value:
        return "网页文件解析失败，请确认文件为有效的 HTML"
    if source_type == DocumentSourceType.EXCEL.value:
        return "表格文件解析失败，请确认为有效的 .xlsx 文件"
    if source_type == DocumentSourceType.EPUB.value:
        return "电子书解析失败，请确认为有效的 EPUB 文件"
    return "资料解析失败，请确认文件格式有效后重试"
