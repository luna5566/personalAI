import zipfile
from pathlib import Path
from uuid import uuid4

import fastapi
import pytest

from app.ai import ocr_provider
from app.api.routes.documents import _detect_source_type
from app.core.config import settings
from app.models.document import Document, DocumentSourceType
from app.services import parsing_service as parsing_service_module
from app.services.parsing_service import (
    SCANNED_PDF_PAGE_MIN_CHARS,
    DocumentParsingError,
    parse_document,
)
from app.storage import storage_service


@pytest.fixture
def storage_root(tmp_path: Path, monkeypatch):
    monkeypatch.setattr(settings, "storage_root", tmp_path)
    return tmp_path


def _document(source_type: DocumentSourceType, filename: str) -> Document:
    return Document(
        title=filename,
        user_id=uuid4(),
        source_type=source_type.value,
        file_path=filename,
        storage_backend="local",
        storage_scope=storage_service.current_storage_scope(),
        mime_type=None,
    )


def _write_and_parse(storage_root: Path, source_type: DocumentSourceType, filename: str, payload: bytes):
    (storage_root / filename).write_bytes(payload)
    return parse_document(_document(source_type, filename))


def test_detects_new_upload_types() -> None:
    assert _detect_source_type("报告.docx", None) == DocumentSourceType.DOCX
    assert _detect_source_type("page.html", "text/html") == DocumentSourceType.HTML
    assert _detect_source_type("table.xlsx", None) == DocumentSourceType.EXCEL
    assert _detect_source_type("book.epub", "application/epub+zip") == DocumentSourceType.EPUB


def test_rejects_unsupported_office_formats() -> None:
    with pytest.raises(fastapi.HTTPException):
        _detect_source_type("legacy.doc", "application/msword")
    with pytest.raises(fastapi.HTTPException):
        _detect_source_type("legacy.xls", "application/vnd.ms-excel")


def test_parse_docx(storage_root: Path) -> None:
    from docx import Document as DocxDocument

    path = storage_root / "note.docx"
    doc = DocxDocument()
    doc.add_paragraph("第一段内容")
    doc.add_paragraph("第二段内容")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "列A"
    table.rows[0].cells[1].text = "列B"
    doc.save(str(path))

    parsed = parse_document(_document(DocumentSourceType.DOCX, "note.docx"))
    assert "第一段内容" in parsed.raw_text
    assert "第二段内容" in parsed.raw_text
    assert "列A\t列B" in parsed.raw_text
    assert parsed.metadata["parser"] == "docx"


def test_parse_html(storage_root: Path) -> None:
    html = (
        "<html><head><style>.x{color:red}</style><script>var a=1;</script></head>"
        "<body><h1>标题</h1><p>第一段&nbsp;文本</p><div>第二段</div></body></html>"
    )
    parsed = _write_and_parse(storage_root, DocumentSourceType.HTML, "page.html", html.encode("utf-8"))
    assert "标题" in parsed.raw_text
    assert "第一段 文本" in parsed.raw_text
    assert "第二段" in parsed.raw_text
    assert "var a=1" not in parsed.raw_text
    assert ".x{color:red}" not in parsed.raw_text
    assert parsed.metadata["parser"] == "html"


def test_parse_excel(storage_root: Path) -> None:
    from openpyxl import Workbook

    path = storage_root / "table.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "数据"
    worksheet.append(["名称", "数量"])
    worksheet.append(["苹果", 3])
    workbook.save(str(path))

    parsed = parse_document(_document(DocumentSourceType.EXCEL, "table.xlsx"))
    assert "## 数据" in parsed.raw_text
    assert "名称\t数量" in parsed.raw_text
    assert "苹果\t3" in parsed.raw_text
    assert parsed.metadata["parser"] == "openpyxl"


def test_parse_epub(storage_root: Path) -> None:
    container = (
        '<?xml version="1.0"?><container version="1.0" '
        'xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
        '<rootfiles><rootfile full-path="OEBPS/content.opf" '
        'media-type="application/oebps-package+xml"/></rootfiles></container>'
    )
    opf = (
        '<?xml version="1.0"?><package xmlns="http://www.idpf.org/2007/opf" version="3.0">'
        "<metadata/><manifest>"
        '<item id="c1" href="ch1.xhtml" media-type="application/xhtml+xml"/>'
        '<item id="c2" href="ch2.xhtml" media-type="application/xhtml+xml"/>'
        "</manifest><spine><itemref idref=\"c1\"/><itemref idref=\"c2\"/></spine></package>"
    )
    ch1 = "<html><body><h2>第一章</h2><p>开头内容</p></body></html>"
    ch2 = "<html><body><h2>第二章</h2><p>结尾内容</p></body></html>"
    path = storage_root / "book.epub"
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("mimetype", "application/epub+zip")
        archive.writestr("META-INF/container.xml", container)
        archive.writestr("OEBPS/content.opf", opf)
        archive.writestr("OEBPS/ch1.xhtml", ch1)
        archive.writestr("OEBPS/ch2.xhtml", ch2)

    parsed = parse_document(_document(DocumentSourceType.EPUB, "book.epub"))
    assert "第一章" in parsed.raw_text
    assert "开头内容" in parsed.raw_text
    assert "第二章" in parsed.raw_text
    assert "结尾内容" in parsed.raw_text
    assert parsed.raw_text.index("第一章") < parsed.raw_text.index("第二章")
    assert parsed.metadata["parser"] == "epub"


def test_scanned_pdf_fails_clearly_when_ocr_disabled(storage_root: Path, monkeypatch) -> None:
    from pypdf import PdfWriter

    path = storage_root / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as handle:
        writer.write(handle)

    monkeypatch.setattr(settings, "ocr_provider", "disabled")
    with pytest.raises(DocumentParsingError) as exc_info:
        parse_document(_document(DocumentSourceType.PDF, "scanned.pdf"))
    assert "扫描版" in str(exc_info.value)


def test_scanned_pdf_uses_ocr_provider(storage_root: Path, monkeypatch) -> None:
    from pypdf import PdfWriter

    path = storage_root / "scanned.pdf"
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    with open(path, "wb") as handle:
        writer.write(handle)

    monkeypatch.setattr(settings, "ocr_provider", "openai_compatible")

    class FakeOCR(ocr_provider.OCRProvider):
        def __init__(self) -> None:
            self.paths: list[Path] = []

        def extract_text(self, path: Path, mime_type: str | None = None) -> str:
            self.paths.append(Path(path))
            return "扫描页识别文本"

    fake = FakeOCR()
    monkeypatch.setattr("app.services.parsing_service.get_ocr_provider", lambda: fake)

    parsed = parse_document(_document(DocumentSourceType.PDF, "scanned.pdf"))
    assert "[第 1 页]" in parsed.raw_text
    assert "扫描页识别文本" in parsed.raw_text
    assert parsed.metadata["parser"] == "pypdf+ocr"
    assert fake.paths and all(p.suffix == ".png" for p in fake.paths)
    assert all(not p.exists() for p in fake.paths)


def test_scanned_pdf_threshold_is_positive() -> None:
    assert SCANNED_PDF_PAGE_MIN_CHARS > 0


def _fake_reader(monkeypatch, page_texts: list[str]) -> None:
    import types

    pages = [
        types.SimpleNamespace(extract_text=lambda text=text: text)
        for text in page_texts
    ]
    monkeypatch.setattr(
        "app.services.parsing_service.PdfReader",
        lambda path: types.SimpleNamespace(pages=pages),
    )


def test_mixed_pdf_ocrs_only_low_text_pages(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_provider", "openai_compatible")
    ocr_pages: list[list[int]] = []

    def fake_ocr_pages(path: Path, page_indexes: list[int]) -> list[str]:
        ocr_pages.append(list(page_indexes))
        return ["OCR 识别文本" for _ in page_indexes]

    monkeypatch.setattr(parsing_service_module, "_ocr_pdf_pages", fake_ocr_pages)
    _fake_reader(monkeypatch, [
        "第一页有足够的正文内容可以入索引",
        "",
        "第三页有足够的正文内容可以入索引",
    ])

    parsed = parsing_service_module._parse_pdf_file(tmp_path / "mixed.pdf")

    assert ocr_pages == [[1]]
    assert "[第 2 页]\nOCR 识别文本" in parsed.raw_text
    assert "第一页有足够的正文内容" in parsed.raw_text
    assert "第三页有足够的正文内容" in parsed.raw_text
    assert parsed.metadata["parser"] == "pypdf+ocr"


def test_mixed_pdf_keeps_text_pages_when_ocr_disabled(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_provider", "disabled")
    _fake_reader(monkeypatch, [
        "第一页有足够的正文内容可以入索引",
        "",
        "第三页有足够的正文内容可以入索引",
    ])

    parsed = parsing_service_module._parse_pdf_file(tmp_path / "mixed.pdf")

    assert "第一页有足够的正文内容" in parsed.raw_text
    assert "第三页有足够的正文内容" in parsed.raw_text
    assert "[第 2 页]" not in parsed.raw_text
    assert parsed.metadata["parser"] == "pypdf"


def test_mixed_pdf_rejects_excessive_ocr_pages(tmp_path: Path, monkeypatch) -> None:
    monkeypatch.setattr(settings, "ocr_provider", "openai_compatible")
    _fake_reader(
        monkeypatch,
        [""] * (parsing_service_module.SCANNED_PDF_OCR_PAGE_LIMIT + 1),
    )

    with pytest.raises(DocumentParsingError) as exc_info:
        parsing_service_module._parse_pdf_file(tmp_path / "huge.pdf")
    assert "上限" in str(exc_info.value)
